"""Build the editable Canva-importable PPTX and matching PDF from shared slide layouts."""
from pathlib import Path
from dataclasses import dataclass, field
import math
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import fitz

ROOT=Path(__file__).resolve().parent
ASSETS=ROOT/'assets'
W,H=13.333333,7.5
C={'bg':'F6F5EF','ink':'213E32','muted':'71816A','lime':'D9EDB4','sage':'E8EFDC','line':'DBE2D3','paper':'FFFEFA','white':'FFFFFF','warm':'C29B72','dark':'142B22'}
FONTS={'head':('Manrope',700),'heavy':('Manrope',800),'medium':('Manrope',500),'body':('DM Sans',400),'bodybold':('DM Sans',700)}
for key,(family,weight) in FONTS.items():
    filename=f'{"manrope" if family=="Manrope" else "dm-sans"}-{weight}.ttf'
    pdfmetrics.registerFont(TTFont(key,str(ASSETS/'fonts'/filename)))

def wrap(text,width,size,font):
    result=[]
    for paragraph in text.split('\n'):
        words=paragraph.split();line=''
        for word in words:
            trial=(line+' '+word).strip()
            if line and pdfmetrics.stringWidth(trial,font,size)>width*72:
                result.append(line);line=word
            else:line=trial
        result.append(line)
    return result

@dataclass
class Slide:
    title:str
    bg:str=C['bg']
    notes:str=''
    items:list=field(default_factory=list)
    def rect(self,x,y,w,h,fill=None,stroke=None,r=.0,lw=1):self.items.append(dict(kind='rect',x=x,y=y,w=w,h=h,fill=fill,stroke=stroke,r=r,lw=lw))
    def ellipse(self,x,y,w,h,fill,stroke=None):self.items.append(dict(kind='ellipse',x=x,y=y,w=w,h=h,fill=fill,stroke=stroke))
    def text(self,text,x,y,w,size=20,font='body',color=None,leading=1.22,align='left'):
        lines=wrap(text,w,size,font)
        h=len(lines)*size*leading/72+.06
        self.items.append(dict(kind='text',x=x,y=y,w=w,h=h,text='\n'.join(lines),size=size,font=font,color=color or C['ink'],leading=leading,align=align))
        return h
    def image(self,path,x,y,w=None,h=None):
        path=ASSETS/path
        iw,ih=Image.open(path).size
        if w is None:w=h*iw/ih
        if h is None:h=w*ih/iw
        self.items.append(dict(kind='image',path=path,x=x,y=y,w=w,h=h))
    def line(self,x,y,w,color=None,lw=1):self.rect(x,y,w,.001,color or C['line'],None,lw=lw)
    def triangle(self,x,y,w,h,fill):self.items.append(dict(kind='triangle',x=x,y=y,w=w,h=h,fill=fill))
    def footer(self,index,dark=False,note=None):
        muted='AEC1A0' if dark else C['muted'];ink=C['lime'] if dark else C['ink']
        self.line(.72,6.99,11.89,'3A5747' if dark else C['line'])
        self.text('rebound.',.75,7.09,1.4,11,'head',ink)
        self.text(note or 'CONGRESSIONAL APP CHALLENGE  /  NC-04',2.1,7.115,9.4,8,'body',muted)
        self.text(f'{index:02}',12.0,7.08,.5,10,'medium',muted,align='right')
    def header(self,label,title,subtitle=None):
        self.text(label.upper(),.75,.48,11.7,10,'bodybold',C['muted'])
        self.text(title,.72,.92,11.9,38,'head',leading=1.13)
        if subtitle:self.text(subtitle,.76,1.71,11.6,16,'body',C['muted'])
    def phone(self,image,x,y,h=5.6):
        w=h*390/844
        self.rect(x-.07,y-.07,w+.14,h+.14,C['ink'],None,.25)
        self.image(image,x,y,w,h)
        return w

slides=[]
s=Slide('Rebound — Small steps. A steadier routine.')
s.rect(8.65,.35,4.15,6.35,C['sage'],None,.35)
s.image('rebound-icon.png',.75,.6,.52)
s.text('rebound.',1.42,.55,5,30,'heavy')
s.text('CONGRESSIONAL APP CHALLENGE',.77,1.62,7,11,'bodybold',C['muted'])
s.text('Small steps.\nA steadier routine.',.72,2.1,8,46,'head',leading=1.14)
s.text('A recovery companion for athletes\nand everyday people.',.78,3.94,7.6,22,'body',C['muted'],leading=1.3)
s.text('Ritvik Saidapur\nSiddharth Senthilkumar',.79,5.12,6.9,17,'bodybold',leading=1.35)
s.text('North Carolina’s 4th District  ·  NC-04',.79,6.02,6.8,12,'body',C['muted'])
s.rect(6.92,5.32,1.37,1.03,C['white'],C['line'],.12);s.image('CAC-logo.png',7.03,5.42,1.15,.82)
s.phone('native-recovery.png',9.42,.56,6.0)
s.text('Native UI preview · example data',9.08,6.68,3.5,8,'body',C['muted'],align='center')
s.footer(1)
s.notes='Ritvik: Introduce yourselves as Ritvik Saidapur and Siddharth Senthilkumar, presenting Rebound for NC-04. Our goal is to make a clinician-cleared recovery routine easier to understand and follow between appointments. Rebound combines a weekly plan, reminders, guided movement practice, and a personal record. The screen shown here is rendered from our current native components using example data, not a real patient or a clinical result. Do not claim that the prototype speeds healing or replaces a clinician.'
slides.append(s)

s=Slide('The problem — Recovery needs more than instructions.')
s.header('The problem','Recovery needs more than instructions.','Between appointments, a patient still needs clarity, confidence, and consistency.')
for i,(label,question,detail) in enumerate([
 ('Clarity','What is my plan today?','A routine should be easy to find and understand.'),
 ('Confidence','Am I moving as intended?','Instructions need a useful visual reference.'),
 ('Consistency','How will I remember?','The next step should fit into an ordinary week.')]):
 x=.75+i*4.02;s.rect(x,2.75,3.79,3.27,C['paper'],C['line'],.23)
 s.text(f'0{i+1}',x+.27,3.02,1.1,38,'medium',C['muted']);s.text(label.upper(),x+.28,3.83,3.1,10,'bodybold',C['muted'])
 s.text(question,x+.28,4.18,3.16,23,'head',leading=1.2);s.text(detail,x+.28,5.18,3.16,16,'body',C['muted'])
s.footer(2)
s.notes='Siddharth: Describe a hypothetical person returning from an injury. They may have instructions but still struggle to remember a routine or understand a movement when they are on their own. These three questions explain the design problem we chose. We are not presenting survey findings, measured adherence statistics, or evidence that the app has improved clinical outcomes. The intended audience includes athletes and everyday people following an individually cleared plan.'
slides.append(s)

s=Slide('The solution — One calm place for the next step.')
s.header('The solution','One calm place for the next step.','Rebound brings the plan, practice, and personal record together.')
s.rect(.75,2.32,8.15,4.36,C['white'],C['line'],.16);s.image('web-overview.png',.84,2.40,7.97,4.21)
for i,(a,b) in enumerate([('Plan','Know what is scheduled.'),('Practice','See the movement and work at your pace.'),('Reflect','Log a session and prepare for a check-in.')]):
 y=2.48+i*1.34;s.text(f'0{i+1}',9.34,y,.65,14,'medium',C['muted']);s.text(a,10.08,y-.05,2.3,23,'head');s.text(b,10.08,y+.43,2.3,15,'body',C['muted'])
s.footer(3,note='Website interface capture · no patient data')
s.notes='Ritvik: Walk through the main loop. First, find the next planned routine. Second, use an illustrated movement guide and a camera estimate when the movement is supported. Third, record the session and how the day feels. The website and iPhone client share the app’s saved routines and guided workflows, while the native iPhone home and Coach library have their own presentation. Browser and device storage are separate; this is not automatic cross-device synchronization.'
slides.append(s)

s=Slide('Weekly goals — A week you can follow.')
s.text('CONSISTENCY, WITH ROOM TO REST',5.0,.65,7.5,10,'bodybold',C['muted'])
s.text('A week you\ncan follow.',4.96,1.12,7.4,43,'head',leading=1.14)
s.phone('native-recovery.png',1.1,.72,5.95)
for y,num,title,body in [
 (2.96,'01','A meaningful weekly goal','Count completed routines against the days in your plan.'),
 (4.04,'02','A timely nudge','Choose reminder days and a time that fits your routine.'),
 (5.12,'03','Rest belongs in the plan','Extra sets do not earn extra goal credit.')]:
 s.ellipse(5.0,y,.43,.43,C['sage']);s.text(num,5.055,y+.08,.32,10,'bodybold');s.text(title,5.7,y-.02,6.5,21,'head');s.text(body,5.7,y+.43,6.2,16,'body',C['muted'])
s.footer(4,note='Native UI preview · example data · progress measures routine completion, not healing')
s.notes='Siddharth: Point to the weekly progress bar above the green recovery-focus card. The example is one of three planned routines complete. A routine earns credit once on its scheduled local date after the run is finished. Extra sets and duplicate saves do not inflate the goal, and dates before a routine was created are excluded. The plan does not demand catch-up exercise on rest days. The reminder shortcut suggests scheduled routine days, but the user still chooses and confirms notification days and time. Existing native reminders are independent and can still alert after completion.'
slides.append(s)

s=Slide('The Coach — Learn the movement. Then practice.')
s.text('THE COACH',.78,.65,7,10,'bodybold',C['muted']);s.text('Learn the movement.\nThen practice.',.73,1.13,8.1,39,'head',leading=1.18)
s.text('Start the routine first.\nExplore the library when you need it.',.8,2.72,6.6,21,'body',C['muted'],leading=1.35)
s.line(.8,3.75,6.7)
s.text('22',.8,4.05,1.8,52,'head');s.text('illustrated movement guides',2.5,4.3,5.1,21,'body')
s.text('Step-by-step teaching\nSlow practice for supported camera movements\nBand and dowel examples with individual clearance',.84,5.18,7.1,16,'body',C['muted'],leading=1.55)
s.phone('native-coach.png',9.32,.62,6.0)
s.footer(5,note='Native UI preview · example data · exercise choice and dose require individual clearance')
s.notes='Ritvik: Show the native Coach screen: a routine action first, followed by a searchable movement library. The catalog contains 22 guides; that is not a claim to cover every injury. Sixteen guides reference previously checked NHS instructions, while six examples retain review-pending labels. Supported camera movements ask for two slow practice cycles. Equipment examples include bands and a light dowel, but the app cannot determine resistance, band anchoring, or a safe clinical range. Do not describe the illustrations as validated biomechanical simulations.'
slides.append(s)

s=Slide('The camera — Guidance where you are looking.',bg=C['ink'])
s.text('ON-DEVICE CAMERA COACH',4.8,.65,7.5,10,'bodybold','B9CEA7');s.text('Guidance where\nyou are looking.',4.75,1.12,7.5,40,'head',C['white'],leading=1.15)
s.phone('camera-ui.png',1.03,.68,5.97)
for y,title,body in [(3.0,'A full-screen view','A large rep counter and simple controls.'),(4.05,'A timely correction cue','An illustrated adjustment appears inside the camera view.'),(5.1,'Pause, breathe, resume','Stillness pauses the session; set breaks support pacing.')]:
 s.ellipse(4.84,y+.05,.12,.12,C['lime']);s.text(title,5.18,y-.04,7,21,'head',C['white']);s.text(body,5.18,y+.4,6.85,16,'body','C5D6B8')
s.footer(6,True,note='Interface test preview · camera feed omitted · estimates are experimental, not clinical assessment')
s.notes='Siddharth: Explain that the coach runs pose estimation locally on camera frames, then checks movement geometry and timing. Poor visibility, framing, lighting, or detected misalignment can stop counting and show an illustrated cue. The full-screen view includes a counter, pause/exit controls, inactivity pause, and optional set breaks. This slide is a UI test capture with the camera feed omitted; synthetic landmarks opened the correction panel. It is not human accuracy evidence. Nine guides offer experimental estimates across eight motion protocols. The other guides provide mirror/manual counting. Band curls use elbow checks; band rows do not have a validated automatic protocol.'
slides.append(s)

s=Slide('Privacy and care — Your camera. Your choice.')
s.header('Privacy and care','Your camera. Your choice.','Normal coaching does not automatically record or upload a clip.')
for i,(label,title,body) in enumerate([
 ('01','Process locally','Pose estimation runs on the device during coaching.'),
 ('02','Review before sharing','A separate short recording is reviewed before the user chooses to share.'),
 ('03','Keep care connected','Clinic workflows support clinician plans and explicitly shared clips.')]):
 x=.75+4.02*i;s.rect(x,2.65,3.78,3.32,C['paper'],C['line'],.24);s.text(label,x+.28,2.96,1,13,'bodybold',C['muted']);s.text(title,x+.28,3.55,3.18,25,'head');s.text(body,x+.28,4.75,3.18,17,'body',C['muted'],leading=1.3)
s.text('Clinic delivery requires configured services and connected accounts. Rebound does not diagnose injuries.',.8,6.34,11.9,13,'body',C['muted'])
s.footer(7)
s.notes='Ritvik: Separate live coaching from recording. Enabling the coach does not automatically save or upload video. The native recorder is a separate action with a review step, and uploads require a configured Clinic service, authentication, and a connected clinician. The current repository includes these workflows, but a live clinical service is not deployed by the prototype update. Symptom follow-up is conservative and rule-based by default. Optional AI explanations need configured services and opt-in; they do not create a diagnostic or injury-specific prescription.'
slides.append(s)

s=Slide('How it works — Built around one shared recovery loop.')
s.header('How it works','Built around one shared recovery loop.','A native iPhone experience, shared guided workflows, and local pose processing.')
for x,title,tech,body in [(.75,'iPhone','EXPO + REACT NATIVE','Recovery home, Coach library, reminders, and native recording.'),(4.77,'Shared app','REACT + TYPESCRIPT','Routine state, check-ins, guided sessions, and web presentation.'),(8.79,'Camera + care','MEDIAPIPE / SUPABASE','Local pose estimates; optional authenticated Clinic services.')]:
 s.rect(x,2.53,3.78,2.28,C['paper'],C['line'],.23);s.text(tech,x+.26,2.83,3.25,10,'bodybold',C['muted']);s.text(title,x+.26,3.23,3.25,25,'head');s.text(body,x+.26,3.83,3.25,15,'body',C['muted'])
s.rect(.75,5.15,11.82,1.35,C['ink'],None,.19)
for i,(a,b) in enumerate([('Camera frame','Local video input'),('Pose landmarks','Joint estimates'),('Movement checks','Visibility + timing'),('Counter or cue','Feedback in the view')]):
 x=1.04+i*2.94;s.text(a,x,5.46,2.64,18,'head',C['white']);s.text(b,x,5.96,2.64,12,'body','BDD1AB')
s.footer(8)
s.notes='Siddharth: Explain the actual architecture. The mobile home and Coach library are React Native. Detailed guides and Clinic workflows use an origin-restricted WebView connected to the web app. A bridge transfers validated recovery snapshots and explicit actions. MediaPipe estimates pose landmarks, and a rule-based tracker checks geometry, visibility, view, and timing. The app does not load a newly trained exercise-video classifier. Supabase integration is optional and must be deployed and configured. AI assistance: OpenAI Codex helped with implementation and debugging. Be ready to explain the system and describe this assistance honestly; verify current district and challenge requirements.'
slides.append(s)

s=Slide('Video demonstration — See Rebound in action.')
s.header('Video demonstration','See Rebound in action.')
s.rect(.76,1.73,8.9,5.00625,C['dark'],None,.12)
s.ellipse(4.68,3.34,1.0,1.0,C['lime']);s.triangle(5.08,3.64,.28,.36,C['ink'])
s.text('ADD YOUR DEMO VIDEO',2.14,4.66,6.14,18,'head',C['white'],align='center')
s.text('16:9 video placeholder',2.14,5.15,6.14,13,'body','B4C7A6',align='center')
s.text('SHOW THE LOOP',10.03,2.04,2.6,10,'bodybold',C['muted'])
for y,num,title,body in [(2.6,'01','Plan','Show the weekly goal and next routine.'),(3.92,'02','Practice','Open a cleared guide and the camera coach.'),(5.24,'03','Reflect','Save a session and show the updated record.')]:
 s.text(num,10.03,y,.5,11,'bodybold',C['muted']);s.text(title,10.61,y-.06,1.95,20,'head');s.text(body,10.05,y+.44,2.49,14,'body',C['muted'])
s.footer(9,note='Suggested demo: 60–90 seconds · replace this placeholder with your own recording in Canva')
s.notes='Both presenters: This slide reserves a 16:9 region for your actual video. In Canva, delete the placeholder elements, upload your recording, place it inside this area, and preview playback. Suggested sequence, not an official duration rule: about 15 seconds on the weekly plan, 35–45 seconds on a supported exercise/camera and correction, then 15–20 seconds saving and viewing progress. Use only individually cleared movements. If footage uses a simulation, label it. Explain tracking limitations and do not claim measured accuracy. Test playback from the actual presenting device, keep an offline backup, and verify the current Congressional App Challenge video instructions separately.'
slides.append(s)

s=Slide('Next steps — A strong prototype. A clear next step.')
s.header('Next steps','A strong prototype. A clear next step.','Keep improving usefulness while measuring what the system can actually do.')
s.rect(.75,2.55,4.0,3.9,C['ink'],None,.24);s.text('ENGINEERING CHECKS',1.05,2.93,3.35,10,'bodybold','BDD1AB');s.text('73',1.02,3.5,3.1,76,'head',C['lime']);s.text('browser checks passed',1.08,4.82,3.22,20,'body',C['white']);s.text('Plus web build, native type-check,\nand iOS JavaScript export.',1.08,5.61,3.25,13,'body','C0D0B2')
for y,num,title,body in [(2.66,'01','Validate on people and iPhones','Measure rep errors and specific form errors against qualified observers.'),(3.98,'02','Evaluate exercise-specific video data','Resolve data rights, labels, and band-assisted movement coverage.'),(5.3,'03','Connect a dependable care workflow','Configure Clinic services, stable HTTPS hosting, and signed native builds.')]:
 s.text(num,5.25,y,.45,11,'bodybold',C['muted']);s.text(title,5.91,y-.06,6.4,21,'head');s.text(body,5.91,y+.43,6.05,15,'body',C['muted'])
s.footer(10,note='Engineering checks are not clinical validation; no human-video accuracy percentage is established')
s.notes='Ritvik: Report the evidence precisely. There are 73 browser cases checked, including the added weekly-goal integration case, plus successful web build, native type-check, and iOS JavaScript export. Those checks include synthetic landmarks and browser streams; they do not establish real-person form accuracy or physical iPhone performance. No new human-video model has been trained or deployed. Dataset permissions and movement mappings remain unresolved. The next steps are qualified human-observer evaluation, physical device checks, and dependable service/hosting setup. Do not call the app clinically validated or fully trained for every exercise.'
slides.append(s)

s=Slide('Closing — A clearer plan. A steadier routine.',bg=C['ink'])
s.image('rebound-icon.png',.81,.68,.68);s.text('rebound.',1.7,.7,6.4,35,'heavy',C['white'])
s.text('A clearer plan.\nA steadier routine.',.76,2.02,11.9,52,'head',C['white'],leading=1.17)
s.text('Helping people understand the next step\nand make room for the routine.',.85,4.24,9.9,23,'body','C8D8BA',leading=1.3)
s.text('Ritvik Saidapur  ·  Siddharth Senthilkumar',.87,5.86,10.2,17,'bodybold',C['white']);s.text('NC-04  /  Congressional App Challenge',.88,6.38,10.3,13,'body','B9CCA9')
s.rect(10.55,5.72,1.7,1.14,C['white'],None,.12);s.image('CAC-logo.png',10.71,5.79,1.38,.984)
s.footer(11,True)
s.notes='Siddharth: Close with the intended benefit: a clearer plan and a steadier routine between appointments. Invite questions about the weekly goal, the coach, privacy, or how the prototype was built. Avoid turning the product’s goal into a claim of proven injury recovery improvement. If asked about AI accuracy, explain the current experimental coverage and evaluation plan honestly. Thank the audience on behalf of Ritvik Saidapur and Siddharth Senthilkumar, NC-04.'
slides.append(s)

s=Slide('Appendix — Sources, identity, and tool credits.')
s.header('Appendix','Sources, identity, and tool credits.')
s.rect(.77,2.05,2.8,2.23,C['white'],C['line'],.16);s.image('CAC-logo.png',.98,2.19,2.38,1.7)
s.text('Congressional App Challenge',.81,4.6,3.07,15,'head');s.text('Official CAC repository asset.\nEvent identification; no endorsement implied.',.82,5.19,3.0,13,'body',C['muted'],leading=1.4)
credits=[('Rebound identity + app evidence','Current app icon, native components, web UI, and repository branch rebound-mvp.'),('Movement and pose references','NHS sitting / strength / flexibility / balance guidance; Google MediaPipe Pose documentation.'),('Logo + typography','CAC: congressionalappchallenge/houseofcode2024 · frontend/src/assets/CAC-logo.svg. Manrope and DM Sans: SIL Open Font License.'),('Tools + status','React, TypeScript, Expo, React Native, MediaPipe, optional Supabase. OpenAI Codex assisted implementation, debugging, and this presentation.')]
for i,(title,body) in enumerate(credits):
 y=2.08+i*1.12;s.text(title,4.16,y,8.24,17,'head');s.text(body,4.17,y+.4,8.12,13,'body',C['muted'],leading=1.27)
s.text('Screens use example data. Camera slide is a UI test preview, not patient footage. Full source links and import instructions accompany the deck.',.83,6.48,11.8,11,'body',C['muted'])
s.footer(12)
s.notes='Reference links: https://github.com/sidsenthilkumar-dev/congressapp/tree/rebound-mvp ; https://github.com/congressionalappchallenge/houseofcode2024/blob/main/frontend/src/assets/CAC-logo.svg ; https://www.nhs.uk/live-well/exercise/sitting-exercises/ ; https://www.nhs.uk/live-well/exercise/strength-exercises/ ; https://www.nhs.uk/live-well/exercise/flexibility-exercises/ ; https://www.nhs.uk/live-well/exercise/balance-exercises/ ; https://chuoling.github.io/mediapipe/solutions/pose.html ; https://docs.expo.dev/ . The CAC website was not accessible from this cloud environment, so this deck does not assert current competition rules or eligibility. Confirm current requirements with the district office. The presenter supplied the names and NC-04; no school, clinician involvement, interviews, patient results, or original student coding claims have been invented. Logos retain their source colors and proportions.'
slides.append(s)

# Shared layout verification catches off-slide shapes before either export.
for number,slide in enumerate(slides,1):
    for item in slide.items:
        if item['x']<-.001 or item['y']<-.001 or item['x']+item['w']>W+.005 or item['y']+item['h']>H+.005:
            raise ValueError(f'Slide {number}: element outside slide: {item}')

prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
prs.core_properties.title='Rebound | Congressional App Challenge | NC-04'
prs.core_properties.author='Ritvik Saidapur; Siddharth Senthilkumar'
prs.core_properties.subject='Recovery companion prototype presentation'
pdf=canvas.Canvas(str(ROOT/'Rebound-Congressional-App-Challenge.pdf'),pagesize=(W*72,H*72))
pdf.setTitle('Rebound — Congressional App Challenge — NC-04');pdf.setAuthor('Ritvik Saidapur; Siddharth Senthilkumar')

def rgb(value):return tuple(int(value[i:i+2],16)/255 for i in (0,2,4))

def draw_ppt(slide,item):
    k=item['kind'];x,y,w,h=[Inches(item[a]) for a in ['x','y','w','h']]
    if k=='image':slide.shapes.add_picture(str(item['path']),x,y,w,h);return
    if k=='text':
        shape=slide.shapes.add_textbox(x,y,w,h);tf=shape.text_frame;tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for index,line in enumerate(item['text'].split('\n')):
            p=tf.paragraphs[0] if index==0 else tf.add_paragraph();p.text=line
            p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(item['size']*item['leading'])
            from pptx.enum.text import PP_ALIGN
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[item['align']]
            p.font.name=FONTS[item['font']][0];p.font.size=Pt(item['size']);p.font.bold=FONTS[item['font']][1]>=700;p.font.color.rgb=RGBColor.from_string(item['color'])
        return
    typ=MSO_SHAPE.OVAL if k=='ellipse' else MSO_SHAPE.ISOSCELES_TRIANGLE if k=='triangle' else MSO_SHAPE.ROUNDED_RECTANGLE if item.get('r') else MSO_SHAPE.RECTANGLE
    shape=slide.shapes.add_shape(typ,x,y,w,h)
    if k=='triangle':shape.rotation=90
    if item.get('r') and k=='rect':shape.adjustments[0]=min(.25,item['r']/min(item['w'],item['h']))
    if item.get('fill'):shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(item['fill'])
    else:shape.fill.background()
    if item.get('stroke'):shape.line.color.rgb=RGBColor.from_string(item['stroke']);shape.line.width=Pt(item.get('lw',1))
    else:shape.line.fill.background()
    if k=='rect' and item.get('fill')==C['dark'] and item['w']>8:shape.name='DEMO VIDEO PLACEHOLDER — replace with your 16:9 recording'

def draw_pdf(item):
    k=item['kind'];x,y,w,h=[item[a]*72 for a in ['x','y','w','h']];bottom=H*72-y-h
    if k=='image':pdf.drawImage(str(item['path']),x,bottom,w,h,mask='auto');return
    if k=='text':
        pdf.setFillColorRGB(*rgb(item['color']));pdf.setFont(item['font'],item['size'])
        for i,line in enumerate(item['text'].split('\n')):
            baseline=H*72-y-item['size']*.85-i*item['size']*item['leading']
            method=pdf.drawCentredString if item['align']=='center' else pdf.drawRightString if item['align']=='right' else pdf.drawString
            xx=x+w/2 if item['align']=='center' else x+w if item['align']=='right' else x
            method(xx,baseline,line)
        return
    fill=bool(item.get('fill'));stroke=bool(item.get('stroke'))
    if fill:pdf.setFillColorRGB(*rgb(item['fill']))
    if stroke:pdf.setStrokeColorRGB(*rgb(item['stroke']));pdf.setLineWidth(item.get('lw',1))
    if k=='ellipse':pdf.ellipse(x,bottom,x+w,bottom+h,fill=int(fill),stroke=int(stroke))
    elif k=='triangle':
        path=pdf.beginPath();path.moveTo(x,bottom);path.lineTo(x+w,bottom+h/2);path.lineTo(x,bottom+h);path.close();pdf.drawPath(path,fill=1,stroke=0)
    elif item.get('r'):pdf.roundRect(x,bottom,w,h,item['r']*72,fill=int(fill),stroke=int(stroke))
    else:pdf.rect(x,bottom,w,h,fill=int(fill),stroke=int(stroke))

for spec in slides:
    page=prs.slides.add_slide(prs.slide_layouts[6]);page.background.fill.solid();page.background.fill.fore_color.rgb=RGBColor.from_string(spec.bg)
    pdf.setFillColorRGB(*rgb(spec.bg));pdf.rect(0,0,W*72,H*72,fill=1,stroke=0)
    for item in spec.items:draw_ppt(page,item);draw_pdf(item)
    page.notes_slide.notes_text_frame.text=spec.notes
    pdf.showPage()
prs.save(ROOT/'Rebound-Congressional-App-Challenge.pptx');pdf.save()
(ROOT/'Speaker-notes.md').write_text('# Rebound — presenter notes\n\n'+'\n\n'.join(f'## {i}. {slide.title}\n\n{slide.notes}' for i,slide in enumerate(slides,1))+'\n')

# Contact sheet and individual previews for visual review and easy sharing.
previews=ROOT/'previews';previews.mkdir(exist_ok=True)
doc=fitz.open(ROOT/'Rebound-Congressional-App-Challenge.pdf')
thumbs=[]
for i,page in enumerate(doc):
    pix=page.get_pixmap(matrix=fitz.Matrix(1.2,1.2),alpha=False);pix.save(previews/f'slide-{i+1:02}.png')
    thumb=Image.open(previews/f'slide-{i+1:02}.png');thumb.thumbnail((480,270));thumbs.append(thumb)
contact=Image.new('RGB',(1464,1128),'#dde3d7')
for i,thumb in enumerate(thumbs):contact.paste(thumb,((i%3)*488+4,(i//3)*282+4))
contact.save(ROOT/'Rebound-slide-preview.png')
print(f'Built {len(slides)} editable slides, PDF, speaker notes, and previews. All elements fit inside the slides.')
